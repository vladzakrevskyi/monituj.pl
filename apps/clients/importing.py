"""Importing clients from a file - a CSV or an Excel export from the
accounting program (Optima, Symfonia, wFirma, inFakt…) or a sheet typed by
hand.

Three steps, nothing saved before the last one:
1. read_table() turns the file into rows of text cells (ClientImport keeps
   them between requests);
2. guess_mapping() finds which column is the name, the email, the phone,
   the NIP and the note - by the headers, or by what the cells look like -
   and the owner can correct it on the preview;
3. check_rows() says per row what will happen (new, already a client,
   repeated in the file, an error), and run_import() does it.

A client is the same as one already there when the email is the same - a
firm may have several contacts under one NIP, so the NIP doesn't merge
anyone. Available on every plan: clients aren't limited, and the same
people could be typed in one by one anyway."""

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction

from apps.audit.models import AuditEvent
from apps.audit.services import AuditService
from apps.clients.models import Client
from apps.common import nip as nip_rules
from apps.common.exceptions import ValidationAppError

MAX_BYTES = 2 * 1024 * 1024
MAX_UNPACKED = 50 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 30
MAX_CELL = 1000

NAME, EMAIL, PHONE, NIP, NOTE = "name", "email", "phone", "nip", "note"
FIELDS = [
    (NAME, "Nazwa"),
    (EMAIL, "E-mail"),
    (PHONE, "Telefon"),
    (NIP, "NIP"),
    (NOTE, "Notatka"),
]
FIELD_LABELS = dict(FIELDS)
LIMITS = {NAME: 255, PHONE: 32, NOTE: MAX_CELL}

# Headers as they come out of Polish programs and sheets, without Polish
# letters, case and punctuation (see _plain). Exact matches first...
HEADERS = {
    NAME: {
        "nazwa",
        "nazwa klienta",
        "nazwa firmy",
        "nazwa kontrahenta",
        "nazwa pelna",
        "pelna nazwa",
        "nazwa skrocona",
        "kontrahent",
        "klient",
        "firma",
        "odbiorca",
        "imie i nazwisko",
        "nazwisko i imie",
        "name",
        "company",
    },
    EMAIL: {"email", "e mail", "adres email", "adres e mail", "mail", "poczta"},
    PHONE: {
        "telefon",
        "tel",
        "nr telefonu",
        "numer telefonu",
        "telefon komorkowy",
        "komorka",
        "tel kom",
        "phone",
        "mobile",
    },
    NIP: {"nip", "nr nip", "numer nip", "nip kontrahenta", "vat id"},
    NOTE: {"uwagi", "notatka", "notatki", "opis", "komentarz", "note", "notes"},
}
# ...then a word of the header starting so ("Nazwa1" in Optima, "E-mail 2").
HEADER_WORDS = [
    (EMAIL, ("mail", "email")),
    (NIP, ("nip",)),
    (PHONE, ("tel", "komork")),
    (NOTE, ("uwag", "notat", "opis", "komentarz")),
    (NAME, ("nazwa", "kontrahent", "firma", "klient")),
]
# Columns about the firm's own people, not the client ("Opiekun klienta",
# "E-mail handlowca") - never guessed.
NOT_THE_CLIENT = ("opiekun", "handlow", "pracownik", "przedstawiciel")

UNREADABLE = (
    "Nie udało się odczytać pliku. Zapisz go jako CSV albo XLSX i spróbuj ponownie."
)


class FileProblem(ValidationAppError):
    pass


# --- reading ----------------------------------------------------------------------


def read_table(uploaded):
    """The file's rows as lists of text cells - trimmed, without empty rows,
    all the same width."""
    if uploaded.size > MAX_BYTES:
        raise FileProblem("Plik jest za duży - maksymalnie 2 MB.")
    name = uploaded.name.lower()
    content = uploaded.read()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        rows = _read_csv(content)
    elif name.endswith(".xls"):
        raise FileProblem(
            "Starszy format Excela (.xls) nie jest obsługiwany. Zapisz plik "
            "jako .xlsx albo .csv."
        )
    else:
        raise FileProblem("Wybierz plik CSV albo XLSX.")

    rows = [
        [_cell(value) for value in row[:MAX_COLUMNS]]
        for row in rows
        if any(str(value or "").strip() for value in row)
    ]
    if not rows:
        raise FileProblem("Plik jest pusty.")
    if len(rows) > MAX_ROWS + 1:
        raise FileProblem(
            f"Plik ma więcej niż {MAX_ROWS} wierszy. Podziel go na mniejsze "
            "i zaimportuj po kolei."
        )
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel keeps a typed-in NIP or phone as a number: 5213017228.0
        value = int(value)
    if hasattr(value, "date") and callable(value.date):
        value = value.date().isoformat()
    return " ".join(str(value).split())[:MAX_CELL]


def _read_csv(content):
    text = _decode(content)
    if "\x00" in text:
        raise FileProblem(UNREADABLE)
    sample = text[:8192]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample else ""
        delimiter = max(";,\t|", key=first.count)
    try:
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as exc:
        raise FileProblem(UNREADABLE) from exc


POLISH_LETTERS = set("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ")


def _decode(content):
    """Excel saves "CSV" in whatever the system uses: UTF-8 (with a BOM in
    "CSV UTF-8"), UTF-16 ("Tekst Unicode"), Windows-1250 on Windows and Mac
    Central European on a Mac. The one-byte ones can't be told apart by
    rules - the one that gives Polish letters wins."""
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return content.decode("utf-16")
        except UnicodeDecodeError as exc:
            raise FileProblem(UNREADABLE) from exc
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    readings = []
    for encoding in ("cp1250", "mac_latin2"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        readings.append((sum(c in POLISH_LETTERS for c in text), text))
    if not readings:
        raise FileProblem(UNREADABLE)
    return max(readings, key=lambda reading: reading[0])[1]


def _read_xlsx(content):
    from openpyxl import load_workbook

    # An .xlsx is a zip: 2 MB of it could unpack to gigabytes.
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            unpacked = sum(entry.file_size for entry in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise FileProblem(UNREADABLE) from exc
    if unpacked > MAX_UNPACKED:
        raise FileProblem(UNREADABLE)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl's own errors, defusedxml's refusals
        raise FileProblem(UNREADABLE) from exc
    try:
        rows = []
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            row = list(row[:MAX_COLUMNS])
            if any(value not in (None, "") for value in row):
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:  # enough to say it's too many
                    break
        return rows
    finally:
        workbook.close()


# --- columns ----------------------------------------------------------------------


def _plain(text):
    text = text.lower().replace("ł", "l")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def field_for_header(header, exact_only=False):
    plain = _plain(header)
    if not plain:
        return ""
    for key, names in HEADERS.items():
        if plain in names:
            return key
    words = plain.split()
    if exact_only or any(w.startswith(NOT_THE_CLIENT) for w in words):
        return ""
    for key, prefixes in HEADER_WORDS:
        if any(w.startswith(prefixes) for w in words):
            return key
    return ""


def has_header(table):
    """The first row names the columns: a header is recognised and the row
    holds no email address."""
    first = table[0]
    return any(field_for_header(cell) for cell in first) and not any(
        "@" in cell for cell in first
    )


def guess_mapping(table, header):
    """Per column the field it fills, or "" - each field at most once."""
    width = len(table[0])
    mapping = [""] * width
    if header:
        # "E-mail" wins over "E-mail 2" wherever it stands.
        for exact_only in (True, False):
            for index, cell in enumerate(table[0]):
                key = field_for_header(cell, exact_only)
                if key and key not in mapping and not mapping[index]:
                    mapping[index] = key
    body = table[1:] if header else table

    def share(index, test):
        values = [row[index] for row in body if row[index]]
        return sum(1 for v in values if test(v)) / len(values) if values else 0

    # What the headers didn't say, the cells may.
    for key, test in (
        (EMAIL, lambda v: "@" in v),
        (NIP, lambda v: nip_rules.is_valid(nip_rules.clean(v))),
    ):
        if key in mapping:
            continue
        free = [i for i in range(width) if not mapping[i]]
        best = max(free, key=lambda i: share(i, test), default=None)
        if best is not None and share(best, test) >= 0.5:
            mapping[best] = key
    if NAME not in mapping and not header:
        # A list without headers: the first text column is most likely names.
        for index in range(width):
            if not mapping[index] and share(index, lambda v: not v.isdigit()) >= 0.5:
                mapping[index] = NAME
                break
    return mapping


def mapping_problem(mapping):
    """Why the import can't go with this mapping, or None."""
    chosen = [key for key in mapping if key]
    for key in set(chosen):
        if chosen.count(key) > 1:
            return f"Pole „{FIELD_LABELS[key]}” jest wybrane dla kilku kolumn."
    if EMAIL not in chosen:
        return "Wskaż kolumnę z adresem e-mail - bez niego nie wyślemy prośby."
    return None


# --- rows -------------------------------------------------------------------------

NEW, EXISTING, REPEATED, ERROR = "new", "existing", "repeated", "error"
STATUSES = (NEW, EXISTING, REPEATED, ERROR)


@dataclass
class Row:
    number: int  # the line in the file, as the owner sees it in Excel
    values: dict
    status: str = NEW
    notes: list = field(default_factory=list)
    client: Client | None = None
    cells: list = field(default_factory=list)  # (label, value) for the preview

    @property
    def skipped(self):
        return self.status in (ERROR, REPEATED)


def _first_email(text):
    """The cell may hold several addresses ("a@x.pl; b@x.pl") - the first
    valid one is the client's."""
    for part in re.split(r"[;,\s]+", text):
        part = part.strip("<>\"'")
        try:
            validate_email(part)
        except ValidationError:
            continue
        return part
    return ""


def check_rows(owner, table, header, mapping):
    """Every row of the file with what the import will do with it."""
    existing = {}
    for client in Client.objects.filter(owner=owner).order_by("pk"):
        existing.setdefault(client.email.lower(), client)
    seen = {}
    rows = []
    start = 2 if header else 1
    for offset, cells in enumerate(table[1:] if header else table):
        values = {key: "" for key, _ in FIELDS}
        for index, key in enumerate(mapping):
            if key:
                values[key] = cells[index]
        row = Row(number=start + offset, values=values)
        rows.append(row)

        raw_email = values[EMAIL]
        email = _first_email(raw_email)
        if not email:
            row.status = ERROR
            row.notes.append(
                f"Nieprawidłowy adres e-mail „{raw_email}”."
                if raw_email
                else "Brak adresu e-mail."
            )
            continue
        if email != raw_email:
            row.notes.append(f"Kilka adresów w komórce - użyjemy {email}.")
        values[EMAIL] = email

        if values[NIP]:
            cleaned = nip_rules.clean(values[NIP])
            if nip_rules.is_valid(cleaned):
                values[NIP] = cleaned
            else:
                row.notes.append(
                    f"Nieprawidłowy NIP „{values[NIP]}” - dodamy klienta bez NIP."
                )
                values[NIP] = ""
        for key, limit in LIMITS.items():
            values[key] = values[key][:limit]
        if not values[NAME]:
            values[NAME] = email.split("@")[0]
            row.notes.append(f"Brak nazwy - nazwiemy klienta „{values[NAME]}”.")

        key = email.lower()
        if key in seen:
            row.status = REPEATED
            row.notes.append(f"Ten sam adres co w wierszu {seen[key]}.")
            continue
        seen[key] = row.number
        if key in existing:
            row.status = EXISTING
            row.client = existing[key]
    return rows


def summary(rows):
    return {status: sum(1 for r in rows if r.status == status) for status in STATUSES}


UPDATED_FIELDS = (NAME, PHONE, NIP, NOTE)


@transaction.atomic
def run_import(owner, rows, update_existing=False, django_request=None):
    """Saves the new clients - and with update_existing fills the existing
    ones with what the file has (an empty cell changes nothing). ->
    {"created", "updated", "skipped"}."""
    new = [
        Client(owner=owner, **{key: row.values[key] for key, _ in FIELDS})
        for row in rows
        if row.status == NEW
    ]
    Client.objects.bulk_create(new)
    updated = 0
    if update_existing:
        for row in rows:
            if row.status != EXISTING:
                continue
            client = row.client
            changed = [
                key
                for key in UPDATED_FIELDS
                if row.values[key] and row.values[key] != getattr(client, key)
            ]
            if not changed:
                continue
            for key in changed:
                setattr(client, key, row.values[key])
            client.save(update_fields=[*changed, "updated_at"])
            updated += 1
    result = {
        "created": len(new),
        "updated": updated,
        "skipped": len(rows) - len(new) - updated,
    }
    AuditService.log(
        AuditEvent.CLIENTS_IMPORTED,
        actor=owner,
        request=django_request,
        metadata=result,
    )
    return result
